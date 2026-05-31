#!/usr/bin/env python3
"""
生成《双人联机打砖块游戏》最终路演PPT
5分钟 5段式：开场(30s) → 演示(2min) → 技术(1.5min) → 分工(30s) → AI(30s)
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ===== 配色方案 =====
BG_DARK = RGBColor(0x1A, 0x1A, 0x2E)       # 深蓝黑背景
ACCENT_BLUE = RGBColor(0x00, 0xD4, 0xFF)    # 青色强调
ACCENT_GREEN = RGBColor(0x00, 0xFF, 0x88)   # 绿色强调
ACCENT_YELLOW = RGBColor(0xFF, 0xD7, 0x00)  # 金色强调
ACCENT_PURPLE = RGBColor(0xBB, 0x86, 0xFC)  # 紫色强调
ACCENT_RED = RGBColor(0xFF, 0x6B, 0x6B)     # 红色强调
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xAA, 0xAA, 0xAA)
BOX_BG = RGBColor(0x25, 0x25, 0x42)         # 卡片背景

prs = Presentation()
prs.slide_width = Inches(13.333)  # 16:9
prs.slide_height = Inches(7.5)


def add_bg(slide):
    """设置幻灯片深色背景"""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BG_DARK


def add_title_bar(slide, title_text, subtitle_text=None):
    """添加顶部标题栏"""
    # 顶部装饰线
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0), prs.slide_width, Inches(0.05)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT_BLUE
    shape.line.fill.background()

    # 标题
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.3), Inches(11), Inches(0.6))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(32)
    p.font.color.rgb = WHITE
    p.font.bold = True

    if subtitle_text:
        txBox2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11), Inches(0.4))
        tf2 = txBox2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle_text
        p2.font.size = Pt(16)
        p2.font.color.rgb = LIGHT_GRAY

    # 分隔线
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.8), Inches(1.25), Inches(11.7), Inches(0.01)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(0x40, 0x40, 0x60)
    line.line.fill.background()


def add_card(slide, left, top, width, height, title, content_lines, accent_color=None):
    """添加卡片样式文本框"""
    if accent_color is None:
        accent_color = ACCENT_BLUE

    # 卡片背景
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = BOX_BG
    shape.line.color.rgb = RGBColor(0x3A, 0x3A, 0x55)
    shape.line.width = Pt(0.5)

    # 左侧色条
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        left + Inches(0.05), top + Inches(0.15), Inches(0.06), Inches(0.3)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent_color
    bar.line.fill.background()

    # 标题
    txBox = slide.shapes.add_textbox(
        left + Inches(0.25), top + Inches(0.1),
        width - Inches(0.5), Inches(0.35)
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(18)
    p.font.color.rgb = accent_color
    p.font.bold = True

    # 内容
    txBox2 = slide.shapes.add_textbox(
        left + Inches(0.25), top + Inches(0.5),
        width - Inches(0.5), height - Inches(0.65)
    )
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    for i, line in enumerate(content_lines):
        if i == 0:
            p = tf2.paragraphs[0]
        else:
            p = tf2.add_paragraph()
        p.text = line
        p.font.size = Pt(13)
        p.font.color.rgb = LIGHT_GRAY
        p.space_after = Pt(4)


def add_bottom_bar(slide, text="github.com/danchenbumin/breakout  |  C++17 + Raylib + ENet"):
    """底部信息栏"""
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.7), Inches(0.3))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(10)
    p.font.color.rgb = RGBColor(0x60, 0x60, 0x80)
    p.alignment = PP_ALIGN.RIGHT


def add_time_badge(slide, left, top, seconds_text):
    """添加演讲时间提示标签"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left, top, Inches(1.0), Inches(0.3)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT_BLUE
    shape.line.fill.background()
    tf = shape.text_frame
    p = tf.paragraphs[0]
    p.text = seconds_text
    p.font.size = Pt(10)
    p.font.color.rgb = BG_DARK
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER


# ===================================================================
# 第1页：封面 — 开场（30秒）
# ===================================================================
slide1 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide1)
add_time_badge(slide1, Inches(0.5), Inches(0.3), "⏱ 30s 开场")

# 大标题
txBox = slide1.shapes.add_textbox(Inches(1.5), Inches(1.3), Inches(10.3), Inches(1.0))
tf = txBox.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "双人联机打砖块游戏"
p.font.size = Pt(52)
p.font.color.rgb = WHITE
p.font.bold = True
p.alignment = PP_ALIGN.CENTER

# 副标题
txBox2 = slide1.shapes.add_textbox(Inches(2), Inches(2.4), Inches(9.3), Inches(0.6))
tf2 = txBox2.text_frame
p2 = tf2.paragraphs[0]
p2.text = "BREAKOUT — 基于C++与Raylib的面向对象游戏开发"
p2.font.size = Pt(22)
p2.font.color.rgb = ACCENT_BLUE
p2.alignment = PP_ALIGN.CENTER

# 一句话亮点
txBox_hl = slide1.shapes.add_textbox(Inches(2.5), Inches(3.1), Inches(8.3), Inches(0.5))
tf_hl = txBox_hl.text_frame
p_hl = tf_hl.paragraphs[0]
p_hl.text = "\"从一行代码到一个可交付作品 — 13周完整C++游戏开发实践\""
p_hl.font.size = Pt(18)
p_hl.font.color.rgb = ACCENT_GREEN
p_hl.font.italic = True
p_hl.alignment = PP_ALIGN.CENTER

# 装饰线
line = slide1.shapes.add_shape(
    MSO_SHAPE.RECTANGLE,
    Inches(5.5), Inches(3.8), Inches(2.3), Inches(0.04)
)
line.fill.solid()
line.fill.fore_color.rgb = ACCENT_GREEN
line.line.fill.background()

# 信息
info_lines = [
    "四川大学 2025级《面向对象程序设计（C++）》课程最终项目",
    "",
    "演讲者：danchenbumin  |  2026年5月",
    "GitHub: https://github.com/danchenbumin/breakout"
]
txBox3 = slide1.shapes.add_textbox(Inches(2), Inches(4.1), Inches(9.3), Inches(1.6))
tf3 = txBox3.text_frame
tf3.word_wrap = True
for i, line in enumerate(info_lines):
    if i == 0:
        p = tf3.paragraphs[0]
    else:
        p = tf3.add_paragraph()
    p.text = line
    p.font.size = Pt(16)
    p.font.color.rgb = LIGHT_GRAY if i > 0 else WHITE
    p.alignment = PP_ALIGN.CENTER

# 底部技术标签
tags = ["C++17", "Raylib", "ENet", "JSON", "CMake", "Multi-thread"]
tag_width = Inches(1.6)
tag_height = Inches(0.35)
start_x = (prs.slide_width - (len(tags) * Inches(1.8))) / 2
for i, tag in enumerate(tags):
    tx = start_x + i * Inches(1.8)
    shape = slide1.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, tx, Inches(5.7), tag_width, tag_height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = BOX_BG
    shape.line.color.rgb = ACCENT_BLUE
    shape.line.width = Pt(0.5)
    tf = shape.text_frame
    tf.paragraphs[0].text = tag
    tf.paragraphs[0].font.size = Pt(12)
    tf.paragraphs[0].font.color.rgb = ACCENT_BLUE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.word_wrap = False

# ===================================================================
# 第2页：核心功能展示 — 演示（2分钟）
# ===================================================================
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide2)
add_title_bar(slide2, "核心功能展示", "六大核心功能 · 现场演示 / 视频播放")
add_time_badge(slide2, Inches(0.5), Inches(0.05), "⏱ 2min 演示")

# 6个功能卡片（2行×3列）
cards_data = [
    ("🎮 双人局域网联机", ACCENT_BLUE,
     ["Host/Client架构，ENet UDP可靠传输",
      "每1/30秒状态同步，双快照线性插值",
      "一人控底部球拍，另一人控顶部"]),
    ("📝 JSON驱动关卡系统", ACCENT_GREEN,
     ["3个难度递增关卡（入门→空心→十字）",
      "球速/布局/颜色全部JSON配置",
      "文件缺失或格式错误自动降级"]),
    ("💾 实时存档/读档", ACCENT_YELLOW,
     ["通关自动保存，退出自动保存",
      "启动检测存档，可选继续或新游戏",
      "存档带版本号，支持未来升级"]),
    ("✏️ 运行时关卡编辑器", ACCENT_PURPLE,
     ["按E进入编辑模式，鼠标添加/删除砖块",
      "网格对齐、重叠检测、边界保护",
      "按S一键保存为custom.json"]),
    ("🎨 粒子特效系统", ACCENT_RED,
     ["砖块破碎/道具拾取粒子反馈",
      "对象池100个粒子预分配，零动态分配",
      "CPU占用降低约30%"]),
    ("🎯 道具系统", RGBColor(0xFF, 0xAA, 0x00),
     ["加长板(W)：球拍宽度+50%",
      "减速球(S)：球速降至70%",
      "额外生命(+)：生命数+1"]),
]

card_w = Inches(3.8)
card_h = Inches(1.7)
start_x = Inches(0.6)
start_y = Inches(1.5)
gap_x = Inches(0.3)
gap_y = Inches(0.2)

for i, (title, color, lines) in enumerate(cards_data):
    row = i // 3
    col = i % 3
    x = start_x + col * (card_w + gap_x)
    y = start_y + row * (card_h + gap_y)
    add_card(slide2, x, y, card_w, card_h, title, lines, color)

# 演示占位提示
demo_y = start_y + 2 * (card_h + gap_y) + Inches(0.15)
demo_box = slide2.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(0.6), demo_y, Inches(12.1), Inches(0.55)
)
demo_box.fill.solid()
demo_box.fill.fore_color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
demo_box.line.color.rgb = ACCENT_BLUE
demo_box.line.width = Pt(1)
tf = demo_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "🎬  此处播放2分钟演示视频（或现场运行游戏） — 展示：单机打砖块 → 双人联机对战 → 道具拾取 → 编辑模式 → 存档读档"
p.font.size = Pt(14)
p.font.color.rgb = ACCENT_BLUE
p.alignment = PP_ALIGN.CENTER
p.font.bold = True

# 操作说明
ops_y = demo_y + Inches(0.7)
txBox = slide2.shapes.add_textbox(Inches(0.8), ops_y, Inches(11.7), Inches(0.25))
tf = txBox.text_frame
p = tf.paragraphs[0]
p.text = "操作：A/D 移动  |  R 重置  |  L 异步加载  |  E 编辑模式  |  ESC 退出  |  道具：W 加长板  S 减速球  + 额外生命"
p.font.size = Pt(11)
p.font.color.rgb = LIGHT_GRAY
p.alignment = PP_ALIGN.CENTER

# ===================================================================
# 第3页：关键技术亮点 — 技术（1.5分钟）
# ===================================================================
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide3)
add_title_bar(slide3, "关键技术亮点", "菱形继承架构 · 三大性能优化 · 13周渐进式迭代")
add_time_badge(slide3, Inches(0.5), Inches(0.05), "⏱ 1.5min 技术")

# 左侧：类继承体系（紧凑版）
add_card(slide3, Inches(0.5), Inches(1.5), Inches(4.0), Inches(2.8),
    "◆ 菱形继承体系", [
        "         GameObject (位置)",
        "        /              \\",
        "  PhysicalObject    VisualObject",
        "  (速度/碰撞)       (颜色/绘制)",
        "        \\              /",
        "   Ball / Paddle / Brick / PowerUp",
        "",
        "· virtual 继承解决二义性",
        "· 纯虚函数 Draw() 实现多态",
        "· unique_ptr 管理对象生命周期",
    ], ACCENT_BLUE)

# 中列：设计模式
add_card(slide3, Inches(4.8), Inches(1.5), Inches(4.0), Inches(2.8),
    "◆ 设计模式应用", [
        "🏭 工厂模式 → PowerUpEffect 抽象基类",
        "🔒 单例模式 → TextureCache 线程安全",
        "♻️ 对象池   → Particle[100] 零分配",
        "🔄 状态模式 → enum GameState 流程控制",
        "🎯 策略模式 → 多态 Apply() 道具效果",
        "",
        "底层：Raylib(渲染) + ENet(网络)",
        "配置：nlohmann/json + CMake 3.10+",
    ], ACCENT_PURPLE)

# 右侧：3个核心技术亮点
tech_cards_data = [
    ("空间网格碰撞", ACCENT_BLUE,
     ["8×6网格划分，每格100×100px",
      "O(N×M) → 接近O(1)",
      "CPU占用：30% → <5%"]),
    ("网络状态插值", ACCENT_GREEN,
     ["双快照(Last/Next)时间戳插值",
      "线性插值消除UDP丢包/乱序抖动",
      "30 tick/s同步，画面流畅"]),
    ("粒子对象池", ACCENT_YELLOW,
     ["预分配Particle[100]静态数组",
      "active标志位管理，O(1)分配",
      "零new/delete，无内存碎片"]),
]

tech_card_h = Inches(0.82)
for i, (title, color, lines) in enumerate(tech_cards_data):
    add_card(slide3, Inches(9.1), Inches(1.5) + i * (tech_card_h + Inches(0.15)),
             Inches(3.7), tech_card_h, title, lines, color)

# 底部：13周时间线（紧凑横排）
timeline_y = Inches(4.6)
txBox_tl = slide3.shapes.add_textbox(Inches(0.5), timeline_y, Inches(2.0), Inches(0.3))
tf_tl = txBox_tl.text_frame
p_tl = tf_tl.paragraphs[0]
p_tl.text = "◆ 13周渐进式迭代开发"
p_tl.font.size = Pt(14)
p_tl.font.color.rgb = ACCENT_GREEN
p_tl.font.bold = True

timeline = [
    ("W1-3\n基础框架", ACCENT_BLUE),
    ("W4-6\nOOP封装", ACCENT_BLUE),
    ("W7\n碰撞检测", ACCENT_BLUE),
    ("W8\n道具+粒子", ACCENT_BLUE),
    ("W9\n双人联机", ACCENT_GREEN),
    ("W10\n多线程", ACCENT_GREEN),
    ("W11\n性能优化", ACCENT_GREEN),
    ("W12\n数据持久化", ACCENT_YELLOW),
    ("W13\n项目收尾", ACCENT_PURPLE),
]

tl_start_x = Inches(0.5)
tl_y = timeline_y + Inches(0.4)
tl_box_w = Inches(1.25)
tl_box_h = Inches(0.85)
tl_gap = Inches(0.1)
tl_arrow_w = Inches(0.15)

for i, (label, color) in enumerate(timeline):
    x = tl_start_x + i * (tl_box_w + tl_gap)
    # 箭头（除最后一个）
    if i < len(timeline) - 1:
        arrow = slide3.shapes.add_shape(
            MSO_SHAPE.RIGHT_ARROW,
            x + tl_box_w + Inches(0.01), tl_y + Inches(0.3),
            tl_gap - Inches(0.02), Inches(0.25)
        )
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = RGBColor(0x40, 0x40, 0x60)
        arrow.line.fill.background()

    # 时间线盒子
    box = slide3.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, x, tl_y, tl_box_w, tl_box_h
    )
    box.fill.solid()
    box.fill.fore_color.rgb = BOX_BG
    box.line.color.rgb = color
    box.line.width = Pt(0.5)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = label
    p.font.size = Pt(10)
    p.font.color.rgb = color
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

# 项目统计
stats_y = tl_y + tl_box_h + Inches(0.15)
stats_box = slide3.shapes.add_textbox(Inches(0.5), stats_y, Inches(12.3), Inches(0.3))
tf = stats_box.text_frame
p = tf.paragraphs[0]
p.text = "📊 12次Git提交 · 2个Pull Request · Feature分支开发 · Conventional Commits规范  |  ~1050行C++ · 11个源文件 · 3个JSON关卡"
p.font.size = Pt(11)
p.font.color.rgb = LIGHT_GRAY
p.alignment = PP_ALIGN.CENTER

# 底部关键挑战（一行）
challenges_y = stats_y + Inches(0.35)
txBox_ch = slide3.shapes.add_textbox(Inches(0.8), challenges_y, Inches(11.7), Inches(0.3))
tf_ch = txBox_ch.text_frame
p_ch = tf_ch.paragraphs[0]
p_ch.text = "关键挑战：菱形继承二义性 → virtual  |  网络抖动 → 双快照插值  |  碰撞O(N×M) → 空间网格  |  内存碎片 → 对象池  |  JSON容错 → try-catch降级"
p_ch.font.size = Pt(10)
p_ch.font.color.rgb = RGBColor(0x80, 0x80, 0xA0)
p_ch.alignment = PP_ALIGN.CENTER

add_bottom_bar(slide3)

# ===================================================================
# 第4页：个人独立开发 — 分工（30秒）
# ===================================================================
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide4)
add_title_bar(slide4, "个人独立开发", "一人完成从设计到交付的全流程 · 13周渐进式迭代")
add_time_badge(slide4, Inches(0.5), Inches(0.05), "⏱ 30s 分工")

# 左侧：角色与职责
add_card(slide4, Inches(0.5), Inches(1.5), Inches(4.0), Inches(3.2),
    "👤 独立承担全部角色", [
        "🏗️ 架构设计 — 菱形继承体系、五大设计模式",
        "💻 程序开发 — ~1050行C++，11个源文件",
        "🎨 游戏设计 — 3个递增难度关卡、3种道具",
        "🌐 网络编程 — ENet Host/Client双人联机",
        "⚡ 性能优化 — 网格碰撞、对象池、异步加载",
        "📝 文档撰写 — README、设计报告、PPT",
        "🔧 工程管理 — Git分支、CMake构建、13周迭代",
    ], ACCENT_GREEN)

# 中列：项目统计
add_card(slide4, Inches(4.8), Inches(1.5), Inches(4.0), Inches(3.2),
    "📊 项目统计", [
        "代码规模：~1050行 C++17",
        "源文件：11个（.cpp/.h）",
        "设计模式：5种（工厂/单例/对象池/状态/策略）",
        "关卡：3个 JSON驱动（入门→空心→十字）",
        "第三方库：Raylib + ENet + nlohmann/json",
        "Git提交：12次，Conventional Commits规范",
        "开发周期：13周渐进式迭代",
    ], ACCENT_BLUE)

# 右侧：核心收获
add_card(slide4, Inches(9.1), Inches(1.5), Inches(3.7), Inches(3.2),
    "✅ 核心收获", [
        "🎯 C++ OOP知识完整实践",
        "  菱形继承 · 虚函数多态",
        "  STL · 智能指针 · 模板",
        "",
        "🎯 工程能力全面提升",
        "  Git · CMake · 设计模式",
        "  代码规范 · 文档撰写",
        "",
        "🎯 实战领域经验",
        "  网络编程(ENet/UDP)",
        "  多线程(std::async)",
        "  文件I/O与JSON序列化",
    ], ACCENT_PURPLE)

# 底部引语
quote_y = Inches(5.0)
txBox_q = slide4.shapes.add_textbox(Inches(2), quote_y, Inches(9.3), Inches(0.5))
tf_q = txBox_q.text_frame
p_q = tf_q.paragraphs[0]
p_q.text = "\"从 'Hello World' 到可交付作品 — 一个人，十三周，一个完整游戏\""
p_q.font.size = Pt(20)
p_q.font.color.rgb = ACCENT_BLUE
p_q.font.italic = True
p_q.alignment = PP_ALIGN.CENTER

add_bottom_bar(slide4)

# ===================================================================
# 第5页：AI 辅助开发 — AI协作（30秒）
# ===================================================================
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide5)
add_title_bar(slide5, "AI 辅助开发", "合理使用AI工具 · 理解为主、AI为辅 · 总占比约10%")
add_time_badge(slide5, Inches(0.5), Inches(0.05), "⏱ 30s AI")

# 三列布局
# 左侧：工具与占比
add_card(slide5, Inches(0.5), Inches(1.5), Inches(3.8), Inches(3.0),
    "🤖 工具与占比", [
        "GitHub Copilot — 代码补全 (~5%)",
        "  重复模板生成、getter/setter补全",
        "",
        "ChatGPT / Claude — 咨询 (~5%)",
        "  架构方案对比、Bug原因分析",
        "",
        "📊 总AI辅助率：约10%",
        "  符合课程评分标准（≤10%）",
        "",
        "✅ 所有AI代码均经人工审核",
    ], ACCENT_BLUE)

# 中间：典型案例
add_card(slide5, Inches(4.6), Inches(1.5), Inches(3.8), Inches(3.0),
    "💡 典型AI辅助案例", [
        "1️⃣ JSON解析框架",
        "   AI生成nlohmann/json使用模板",
        "   人工添加错误降级+默认配置",
        "",
        "2️⃣ CMakeLists.txt构建配置",
        "   AI生成基础CMake模板",
        "   人工添加ENet查找+文件复制",
        "",
        "3️⃣ 菱形继承Bug定位",
        "   AI分析二义性原因",
        "   人工决定virtual继承方案",
    ], ACCENT_GREEN)

# 右侧：使用原则
add_card(slide5, Inches(8.7), Inches(1.5), Inches(3.8), Inches(3.0),
    "📋 AI使用原则", [
        "✅ AI是效率工具，不是代码替代品",
        "",
        "✅ 所有AI生成代码经过人工审核",
        "",
        "✅ 理解代码原理，能独立解释",
        "",
        "✅ 关键AI辅助部分注释标注来源",
        "",
        "✅ AI提供新思路，人工做最终决策",
        "",
        "✅ 学会判断AI输出是否正确、最优",
    ], ACCENT_PURPLE)

# 底部总结
summary_y = Inches(4.8)
txBox_s = slide5.shapes.add_textbox(Inches(0.6), summary_y, Inches(12), Inches(0.8))
tf_s = txBox_s.text_frame
tf_s.word_wrap = True
p_s = tf_s.paragraphs[0]
p_s.text = '核心观点：AI改变了开发方式，但没有改变“理解代码才能写出好代码”的本质。'
p_s.font.size = Pt(18)
p_s.font.color.rgb = WHITE
p_s.alignment = PP_ALIGN.CENTER
p_s.font.bold = True
p_s2 = tf_s.add_paragraph()
p_s2.text = "在课程中使用AI，关键是学会判断AI的输出是否正确、是否最优 — 这比不用AI更需要理解代码。"
p_s2.font.size = Pt(14)
p_s2.font.color.rgb = LIGHT_GRAY
p_s2.alignment = PP_ALIGN.CENTER

add_bottom_bar(slide5)

# ===================================================================
# 第6页：感谢聆听 · Q&A
# ===================================================================
slide6 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide6)

# 大标题
txBox = slide6.shapes.add_textbox(Inches(1), Inches(1.0), Inches(11.3), Inches(1.0))
tf = txBox.text_frame
p = tf.paragraphs[0]
p.text = "感谢聆听 · 欢迎提问"
p.font.size = Pt(52)
p.font.color.rgb = WHITE
p.font.bold = True
p.alignment = PP_ALIGN.CENTER

# 装饰
line = slide6.shapes.add_shape(
    MSO_SHAPE.RECTANGLE,
    Inches(5.5), Inches(2.0), Inches(2.3), Inches(0.04)
)
line.fill.solid()
line.fill.fore_color.rgb = ACCENT_GREEN
line.line.fill.background()

# 左侧：准备回答的问题
add_card(slide6, Inches(0.5), Inches(2.4), Inches(6.2), Inches(2.6),
    "💡 准备好回答的问题", [
        "1. 为什么选择菱形继承而不是组合模式？",
        "   → 展示C++多继承特性；实际项目中组合更灵活",
        "",
        "2. 网络同步为什么用UDP(ENet)而不是TCP？",
        "   → 游戏实时性要求低延迟，ENet在UDP上实现可靠传输",
        "",
        "3. 关卡编辑器如何保证保存数据的合法性？",
        "   → 网格对齐 + 位置验证 + 边界检查",
        "",
        "4. 如果重新设计，会做哪些架构调整？",
        "   → ECS架构替代菱形继承，更好的解耦和扩展性",
    ], ACCENT_BLUE)

# 右侧：待改进与未来（原第6页内容精简）
add_card(slide6, Inches(7.0), Inches(2.4), Inches(5.8), Inches(2.6),
    "⚠️ 待改进 · 🚀 未来计划", [
        "⚠️ 音效系统缺失 — 无碰撞音效/背景音乐",
        "⚠️ UI较为简陋 — 缺少图形化菜单界面",
        "⚠️ 仅支持局域网 — 需扩展到互联网联机",
        "",
        "🚀 移植移动端（Android/iOS）",
        "🚀 发布到 Steam / itch.io",
        "🚀 迁移 OpenGL 实现 3D 渲染",
        "🚀 学习 ECS架构 · C++20/23新特性",
        "🚀 参加 GitHub Game Off 比赛",
    ], ACCENT_YELLOW)

# 底部信息
txBox2 = slide6.shapes.add_textbox(Inches(2), Inches(5.4), Inches(9.3), Inches(1.0))
tf2 = txBox2.text_frame
tf2.word_wrap = True
p = tf2.paragraphs[0]
p.text = "GitHub: https://github.com/danchenbumin/breakout"
p.font.size = Pt(18)
p.font.color.rgb = ACCENT_BLUE
p.alignment = PP_ALIGN.CENTER
p.font.bold = True
p2 = tf2.add_paragraph()
p2.text = "可现场运行演示 · 可展示代码仓库 · 欢迎技术交流"
p2.font.size = Pt(14)
p2.font.color.rgb = LIGHT_GRAY
p2.alignment = PP_ALIGN.CENTER

# ===== 保存 =====
output_path = r"d:\Linux c++ class code\ubuntu\breakout\双人联机打砖块游戏_最终路演.pptx"
prs.save(output_path)
print(f"PPT已保存到: {output_path}")
print(f"共 {len(prs.slides)} 张幻灯片")
print("5段式结构：开场(30s) → 演示(2min) → 技术(1.5min) → 分工(30s) → AI(30s)")
